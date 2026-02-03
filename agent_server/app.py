#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import csv
import io
import os
import threading
from dataclasses import dataclass
from typing import Any, Dict, List, Tuple, Optional

import requests
import tkinter as tk
from tkinter import ttk, filedialog, messagebox


BACKEND_URL = os.getenv("AGENT_BACKEND_URL", "http://agent_server:8000")


def _parse_list(text: str) -> List[str]:
    return [item.strip() for item in text.split(",") if item.strip()]


def _load_preview_from_csv(data: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = [row for row in reader]
    return rows[:5], list(reader.fieldnames or [])


def _load_preview_from_xlsx(data: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [], []

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    preview: List[Dict[str, Any]] = []
    for row in rows[1:6]:
        preview.append({headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))})
    return preview, headers


def _preview_from_bytes(data: bytes, filename: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    ext = os.path.splitext(filename)[1].lower()
    if ext in (".csv", ""):
        return _load_preview_from_csv(data)
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return _load_preview_from_xlsx(data)
    return [], []


def _filename_from_content_disposition(disposition: str, fallback: str) -> str:
    # простой разбор: filename="..."
    if not disposition:
        return fallback
    lower = disposition.lower()
    if "filename=" not in lower:
        return fallback
    part = disposition.split("filename=")[-1].strip().strip("\"' ")
    return part or fallback


@dataclass
class ProcessResult:
    content: bytes
    filename: str
    mime: str


class App(ttk.Frame):
    def __init__(self, master: tk.Tk):
        super().__init__(master)
        self.master = master

        self.selected_path: Optional[str] = None
        self.selected_name: Optional[str] = None

        self.last_result: Optional[ProcessResult] = None

        self._build_ui()

    def _build_ui(self):
        self.master.title("Stroki Agent (Tkinter)")
        self.master.geometry("900x600")

        # Верхняя панель
        top = ttk.LabelFrame(self, text="Входные данные")
        top.pack(fill="x", padx=12, pady=10)

        row1 = ttk.Frame(top)
        row1.pack(fill="x", padx=10, pady=8)

        self.file_label_var = tk.StringVar(value="Файл не выбран")
        ttk.Button(row1, text="Выбрать файл…", command=self.on_pick_file).pack(side="left")
        ttk.Label(row1, textvariable=self.file_label_var).pack(side="left", padx=10)

        row2 = ttk.Frame(top)
        row2.pack(fill="x", padx=10, pady=6)

        ttk.Label(row2, text="Inputs (заполненные столбцы):").grid(row=0, column=0, sticky="w")
        self.inputs_var = tk.StringVar()
        ttk.Entry(row2, textvariable=self.inputs_var, width=80).grid(row=0, column=1, sticky="we", padx=8)

        row2.columnconfigure(1, weight=1)

        row3 = ttk.Frame(top)
        row3.pack(fill="x", padx=10, pady=6)

        ttk.Label(row3, text="Tasks (нужно заполнить):").grid(row=0, column=0, sticky="w")
        self.tasks_var = tk.StringVar()
        ttk.Entry(row3, textvariable=self.tasks_var, width=80).grid(row=0, column=1, sticky="we", padx=8)

        row3.columnconfigure(1, weight=1)

        row_images = ttk.Frame(top)
        row_images.pack(fill="x", padx=10, pady=6)

        ttk.Label(row_images, text="Image search column:").grid(row=0, column=0, sticky="w")
        self.image_search_var = tk.StringVar()
        ttk.Entry(row_images, textvariable=self.image_search_var, width=80).grid(row=0, column=1, sticky="we", padx=8)

        row_images.columnconfigure(1, weight=1)

        row_images_count = ttk.Frame(top)
        row_images_count.pack(fill="x", padx=10, pady=6)

        ttk.Label(row_images_count, text="Image count (0-5) or column:").grid(row=0, column=0, sticky="w")
        self.image_count_var = tk.StringVar()
        ttk.Entry(row_images_count, textvariable=self.image_count_var, width=80).grid(row=0, column=1, sticky="we", padx=8)

        row_images_count.columnconfigure(1, weight=1)

        row4 = ttk.Frame(top)
        row4.pack(fill="x", padx=10, pady=10)

        self.process_btn = ttk.Button(row4, text="Обработать", command=self.on_process)
        self.process_btn.pack(side="left")

        self.save_btn = ttk.Button(row4, text="Сохранить результат…", command=self.on_save, state="disabled")
        self.save_btn.pack(side="left", padx=8)

        self.status_var = tk.StringVar(value=f"Backend: {BACKEND_URL}")
        ttk.Label(row4, textvariable=self.status_var).pack(side="left", padx=12)

        # Спиннер (indeterminate progressbar)
        self.pb = ttk.Progressbar(row4, mode="indeterminate", length=180)
        self.pb.pack(side="right")

        # Превью
        preview_box = ttk.LabelFrame(self, text="Превью первых 5 строк результата")
        preview_box.pack(fill="both", expand=True, padx=12, pady=10)

        self.tree = ttk.Treeview(preview_box, show="headings")
        self.tree.pack(side="left", fill="both", expand=True)

        vsb = ttk.Scrollbar(preview_box, orient="vertical", command=self.tree.yview)
        vsb.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=vsb.set)

        self.pack(fill="both", expand=True)

    def on_pick_file(self):
        path = filedialog.askopenfilename(
            title="Выберите CSV/XLSX",
            filetypes=[
                ("CSV files", "*.csv"),
                ("Excel files", "*.xlsx *.xlsm *.xltx *.xltm"),
                ("All files", "*.*"),
            ],
        )
        if not path:
            return

        self.selected_path = path
        self.selected_name = os.path.basename(path)
        self.file_label_var.set(self.selected_name)
        self.status_var.set("Файл выбран. Заполните inputs/tasks и нажмите «Обработать».")

    def _set_busy(self, busy: bool, text: str = ""):
        # progressbar start/stop :contentReference[oaicite:4]{index=4}
        if busy:
            self.process_btn.configure(state="disabled")
            self.save_btn.configure(state="disabled")
            self.pb.start(10)
        else:
            self.pb.stop()
            self.process_btn.configure(state="normal")
            self.save_btn.configure(state="normal" if self.last_result else "disabled")

        if text:
            self.status_var.set(text)

    def on_process(self):
        if not self.selected_path or not os.path.exists(self.selected_path):
            messagebox.showwarning("Нет файла", "Выберите файл CSV/XLSX.")
            return

        inputs_list = _parse_list(self.inputs_var.get())
        tasks_list = _parse_list(self.tasks_var.get())
        image_search_column = self.image_search_var.get().strip()
        image_count_column = self.image_count_var.get().strip()

        if not inputs_list:
            messagebox.showwarning("Inputs", "Укажите inputs через запятую.")
            return
        if not tasks_list:
            messagebox.showwarning("Tasks", "Укажите tasks через запятую.")
            return

        self.last_result = None
        self._clear_preview()
        self._set_busy(True, "Отправляем файл в backend и ждём ответа...")

        # Запрос в отдельном потоке, а UI обновляем через after (Tkinter не thread-safe) :contentReference[oaicite:5]{index=5}
        thread = threading.Thread(
            target=self._process_worker,
            args=(self.selected_path, inputs_list, tasks_list, image_search_column, image_count_column),
            daemon=True,
        )
        thread.start()

    def _process_worker(
        self,
        path: str,
        inputs_list: List[str],
        tasks_list: List[str],
        image_search_column: str,
        image_count_column: str,
    ):
        try:
            with open(path, "rb") as f:
                file_bytes = f.read()

            files = {
                "file": (os.path.basename(path), file_bytes, "application/octet-stream"),
            }
            data = {
                "inputs": ", ".join(inputs_list),
                "tasks": ", ".join(tasks_list),
                "image_search_column": image_search_column,
                "image_count_column": image_count_column,
            }

            resp = requests.post(BACKEND_URL, files=files, data=data, timeout=None)

            if resp.status_code >= 400:
                raise RuntimeError(f"Backend ответил ошибкой {resp.status_code}: {resp.text}")

            content = resp.content
            filename = _filename_from_content_disposition(
                resp.headers.get("Content-Disposition", ""),
                os.path.basename(path),
            )
            mime = resp.headers.get("Content-Type", "application/octet-stream")

            result = ProcessResult(content=content, filename=filename, mime=mime)

            # Вернёмся в UI-поток
            self.master.after(0, lambda: self._on_process_success(result))

        except Exception as exc:
            self.master.after(0, lambda exc=exc: self._on_process_error(exc))

    def _on_process_success(self, result: ProcessResult):
        self.last_result = result
        self._set_busy(False, "Файл обработан. Можно сохранить результат или посмотреть превью.")

        preview_rows, preview_headers = _preview_from_bytes(result.content, result.filename)
        if preview_rows and preview_headers:
            self._populate_preview(preview_rows, preview_headers)
        else:
            # если не получилось превью (например xlsx без openpyxl)
            if os.path.splitext(result.filename)[1].lower() in (".xlsx", ".xlsm", ".xltx", ".xltm"):
                messagebox.showinfo(
                    "Превью недоступно",
                    "Не удалось построить превью. Для XLSX установите openpyxl:\n\npip install openpyxl",
                )
            else:
                messagebox.showinfo("Превью недоступно", "Не удалось построить превью.")

        self.save_btn.configure(state="normal")

    def _on_process_error(self, exc: Exception):
        self._set_busy(False, "Ошибка при обработке. См. сообщение.")
        messagebox.showerror("Ошибка", f"Ошибка запроса к backend:\n\n{exc}")

    def on_save(self):
        if not self.last_result:
            messagebox.showinfo("Нет результата", "Сначала обработайте файл.")
            return

        # Save As dialog :contentReference[oaicite:6]{index=6}
        save_path = filedialog.asksaveasfilename(
            title="Сохранить обработанный файл",
            initialfile=self.last_result.filename,
            defaultextension=os.path.splitext(self.last_result.filename)[1] or ".csv",
            filetypes=[
                ("CSV files", "*.csv"),
                ("Excel files", "*.xlsx *.xlsm *.xltx *.xltm"),
                ("All files", "*.*"),
            ],
        )
        if not save_path:
            return

        try:
            with open(save_path, "wb") as f:
                f.write(self.last_result.content)
            self.status_var.set(f"Сохранено: {save_path}")
        except Exception as exc:
            messagebox.showerror("Ошибка сохранения", str(exc))

    def _clear_preview(self):
        for col in self.tree["columns"]:
            self.tree.heading(col, text="")
        self.tree["columns"] = ()
        for item in self.tree.get_children():
            self.tree.delete(item)

    def _populate_preview(self, rows: List[Dict[str, Any]], headers: List[str]):
        self._clear_preview()

        self.tree["columns"] = headers
        for h in headers:
            self.tree.heading(h, text=h)
            self.tree.column(h, width=150, stretch=True)

        for row in rows:
            values = [row.get(h, "") for h in headers]
            self.tree.insert("", "end", values=values)


def main():
    root = tk.Tk()
    # ttk theme можно поменять, если хочешь
    app = App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
