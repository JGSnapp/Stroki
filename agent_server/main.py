import asyncio
import csv
import hashlib
import hmac
import io
import json
import logging
import os
import time
import urllib.parse
import uuid
from threading import Lock
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, TypedDict

import psycopg2

from fastapi import FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from langchain_core.messages import AIMessage, AnyMessage, HumanMessage, SystemMessage, ToolMessage
from langchain_core.tools import tool
from langchain_openai import ChatOpenAI
from langgraph.graph import END, StateGraph
from pydantic import BaseModel, Field
from pydantic.v1 import BaseModel as BaseModelV1, Field as FieldV1

from watermark import has_watermark_url
from xml_images import xml_images
from xml_search import xml_search

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover - optional dependency
    def load_dotenv(*_args, **_kwargs):
        return None


def _load_env() -> None:
    candidates = [
        Path(__file__).resolve().parent.parent / ".env",
        Path(__file__).resolve().parent / ".env",
    ]
    for candidate in candidates:
        if candidate.is_file():
            load_dotenv(candidate)
            return
    load_dotenv()


_load_env()


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)


MAX_ITERATIONS = int(os.getenv("MAX_ITERATIONS", os.getenv("MAX_ITTERATIONS", "6")))
PAGES_PER_TIME = int(os.getenv("PAGES_PER_TIME", "2"))
MAX_ROWS = 300
MAX_CONCURRENCY = int(os.getenv("MAX_CONCURRENCY", "4"))
LLM_RETRIES = int(os.getenv("LLM_RETRIES", "3"))
LLM_RETRY_WAIT = float(os.getenv("LLM_RETRY_WAIT", "2.0"))
RECURSION_LIMIT = int(os.getenv("RECURSION_LIMIT", "5000"))
DATABASE_URL = os.getenv("DATABASE_URL", "postgresql://stroky:stroky@postgres:5432/stroky")
UPLOAD_DIR = os.getenv("UPLOAD_DIR", "uploads")
IP_HASH_SALT = os.getenv("IP_HASH_SALT", "stroky_salt")

AGENT_PROMPT = (
    "Ты агент, отвечающий на вопросы по материалам из интернета. "
    "У тебя есть возможность использовать два инструмента: "
    "return_answers_tool позволяет тебе дать ответы на вопросы; "
    "ask_more_pages позволяет тебе запросить больше информации "
    "(используй только если в уже предоставленной тебе информации точно нет ответов). "
    "Всегда возвращай return_answers_tool, даже если удалось заполнить только часть полей tasks."
)

SELECTION_PROMPT = (
    "You are given a user query and a list of table column names. "
    "Select up to 3 input columns (inputs) from the existing columns that will be used as context for each row. "
    "Select up to 3 task columns (tasks) that should be filled. Task columns may be new column names derived from the query. "
    "If the user asks for images/photos/pictures, you MUST set img_task and img_count. "
    "img_task must be one existing column to search images for (pick the best item/name column). "
    "img_count must be an integer from 1 to 5 (use 1 if the user did not specify a number). "
    "If images are not needed, set img_task and img_count to null. "
    "If no enrichment is needed, leave inputs and tasks empty."
)

REQUEST_PROMPT = (
    "Тебе даны наименования полей inputs и tasks. "
    "Тебе нужно на основе полей, которые будут передаваться в inputs сформировать запрос, "
    "который найдет в интернете всю необходимую информацию для заполнения полей tasks. "
    "Для обозначения полей inputs используй '*input_1*', '*input_2*' и т.д. "
    "Пример: тебе передан inputs: Наименование товара, а tasks: производитель, страна производства. "
    "Запросы, которые ты должен вывести должны выглядеть примерно так: "
    "'производитель *input_1*' и 'страна производства *input_1*'."
    "Return no more than 3 requests."
)

LAST_PROMPT = (
    "У тебя есть последняя возможность выдать какой-то результат. "
    "Если ты ничего не нашел, ничего не заполняй."
)

SAVE_LOGS = os.getenv("SAVE_LOGS", "false").lower() in ("1", "true", "yes")
LOGS_DIR = Path(__file__).resolve().parent.parent / "logs"
PROGRESS: Dict[str, Dict[str, Any]] = {}
PROGRESS_LOCK = Lock()
RUN_STATS: Dict[str, Dict[str, Any]] = {}
RUN_STATS_LOCK = Lock()
XML_SEARCH_DOCS = int(os.getenv("XML_SEARCH_DOCS", "10"))
XML_SEARCH_MAXPASSAGES = int(os.getenv("XML_SEARCH_MAXPASSAGES", "5"))
XML_SEARCH_GROUPBY = os.getenv("XML_SEARCH_GROUPBY", "10")
XML_SEARCH_LR = os.getenv("XML_SEARCH_LR")
XML_IMAGES_LIMIT = int(os.getenv("XML_IMAGES_LIMIT", "50"))
XML_IMAGES_PAGE = int(os.getenv("XML_IMAGES_PAGE", "0"))
XML_IMAGES_DEVICE = os.getenv("XML_IMAGES_DEVICE")
XML_IMAGES_DOMAIN = os.getenv("XML_IMAGES_DOMAIN")
XML_IMAGES_LR = os.getenv("XML_IMAGES_LR")

os.makedirs(UPLOAD_DIR, exist_ok=True)


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _get_db_conn() -> psycopg2.extensions.connection:
    conn = psycopg2.connect(DATABASE_URL)
    conn.autocommit = True
    return conn


def _hash_ip(ip: str) -> str:
    return hashlib.sha256(f"{IP_HASH_SALT}:{ip}".encode("utf-8")).hexdigest()


def _get_client_ip(request: Request) -> str:
    forwarded = request.headers.get("x-forwarded-for", "")
    if forwarded:
        return forwarded.split(",")[0].strip()
    if request.client and request.client.host:
        return request.client.host
    return "unknown"


def _get_or_create_visitor(ip_hash: str) -> str:
    now = _utc_now()
    with _get_db_conn() as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO visitors (ip_hash, first_seen_at, last_seen_at)
            VALUES (%s, %s, %s)
            ON CONFLICT (ip_hash)
            DO UPDATE SET last_seen_at = EXCLUDED.last_seen_at
            RETURNING id
            """,
            (ip_hash, now, now),
        )
        row = cur.fetchone()
        return str(row[0])


def _create_session(visitor_id: str, user_agent: Optional[str], referrer: Optional[str]) -> str:
    now = _utc_now()
    with _get_db_conn() as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO sessions (visitor_id, started_at, user_agent, referrer)
            VALUES (%s, %s, %s, %s)
            RETURNING id
            """,
            (visitor_id, now, user_agent, referrer),
        )
        row = cur.fetchone()
        return str(row[0])


def _end_session(session_id: str) -> None:
    now = _utc_now()
    with _get_db_conn() as conn, conn.cursor() as cur:
        cur.execute(
            "UPDATE sessions SET ended_at = %s WHERE id = %s AND ended_at IS NULL",
            (now, session_id),
        )


def _store_upload(
    session_id: str,
    original_name: str,
    content_type: str,
    data: bytes,
) -> str:
    file_hash = hashlib.sha256(data).hexdigest()
    ext = os.path.splitext(original_name)[1].lower() or ".bin"
    date_dir = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    stored_dir = os.path.join(UPLOAD_DIR, date_dir)
    os.makedirs(stored_dir, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}{ext}"
    stored_path = os.path.join(stored_dir, stored_name)

    with open(stored_path, "wb") as handle:
        handle.write(data)

    relative_name = os.path.join(date_dir, stored_name).replace("\\", "/")
    with _get_db_conn() as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO uploads (session_id, original_name, stored_name, mime, size_bytes, sha256, created_at)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (session_id, original_name, relative_name, content_type, len(data), file_hash, _utc_now()),
        )
    return relative_name



class API:
    llm: ChatOpenAI

    def __init__(self) -> None:
        proxy_model = os.getenv("PROXY_MODEL") or os.getenv("API_MODEL")
        proxy_api_key = os.getenv("PROXY_API_KEY") or os.getenv("API_KEY")
        proxy_base_url = os.getenv("PROXY_BASE_URL") or os.getenv("BASE_URL", "https://api.openai.com/v1")

        llm_max_tokens = int(os.getenv("LLM_MAX_TOKENS", "1024"))

        logger.info(
            "LLM init: proxy_model=%r, proxy_base_url=%r, proxy_key_set=%s, llm_max_tokens=%s",
            proxy_model,
            proxy_base_url,
            bool(proxy_api_key),
            llm_max_tokens,
        )

        if not proxy_model:
            raise RuntimeError("PROXY_MODEL is not set")
        self.llm = ChatOpenAI(
            model=proxy_model,
            api_key=proxy_api_key,
            base_url=proxy_base_url,
            temperature=0,
            max_tokens=llm_max_tokens,
        )


api = API()
llm = api.llm


def _save_xml_search_log(query: str, results: List[Dict[str, Any]]) -> None:
    if not SAVE_LOGS:
        return
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_query = _safe_filename(query)
    path = LOGS_DIR / "xml_search" / f"{safe_query}-{timestamp}.json"
    os.makedirs(path.parent, exist_ok=True)
    content = {
        "query": query,
        "timestamp": timestamp,
        "results": results,
    }
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(content, handle, ensure_ascii=False, indent=2)
    logger.info("xml_search: saved log=%s", path)


def _save_image_search_log(query: str, urls: List[str]) -> None:
    if not SAVE_LOGS:
        return
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_query = _safe_filename(query)
    path = LOGS_DIR / "xml_images" / f"{safe_query}-{timestamp}.json"
    os.makedirs(path.parent, exist_ok=True)
    content = {
        "query": query,
        "timestamp": timestamp,
        "urls": urls,
    }
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(content, handle, ensure_ascii=False, indent=2)
    logger.info("xml_images: saved log=%s", path)


def _save_start_agent_log(query: str, columns: List[str], selection: "ColumnSelection") -> None:
    if not SAVE_LOGS:
        return
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = LOGS_DIR / f"start_agent_{timestamp}.txt"
    os.makedirs(path.parent, exist_ok=True)
    content = [
        f"timestamp: {timestamp}",
        f"query: {query}",
        f"columns: {', '.join(columns)}",
        f"inputs: {', '.join(selection.inputs)}",
        f"tasks: {', '.join(selection.tasks)}",
        f"img_task: {selection.img_task or ''}",
        f"img_count: {selection.img_count if selection.img_count is not None else ''}",
    ]
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(content) + "\n")
    logger.info("start_agent: saved log=%s", path)


def _init_run_stats(run_id: str) -> None:
    if not SAVE_LOGS:
        return
    with RUN_STATS_LOCK:
        RUN_STATS[run_id] = {
            "llm_count": 0,
            "search_count": 0,
            "image_search_count": 0,
        }


def _inc_run_stat(run_id: Optional[str], key: str, delta: int = 1) -> None:
    if not SAVE_LOGS or not run_id:
        return
    with RUN_STATS_LOCK:
        stats = RUN_STATS.setdefault(
            run_id, {"llm_count": 0, "search_count": 0, "image_search_count": 0}
        )
        stats[key] = int(stats.get(key, 0)) + delta


def _write_run_stats(run_id: str) -> None:
    if not SAVE_LOGS:
        return
    timestamp = run_id
    path = LOGS_DIR / f"count_{timestamp}.txt"
    os.makedirs(path.parent, exist_ok=True)
    with RUN_STATS_LOCK:
        stats = RUN_STATS.get(run_id, {})
        llm_count = int(stats.get("llm_count", 0))
        search_count = int(stats.get("search_count", 0))
        image_search_count = int(stats.get("image_search_count", 0))
    lines = [
        f"llm_count: {llm_count}",
        f"search_count: {search_count}",
        f"image_search_count: {image_search_count}",
    ]
    with open(path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(lines) + "\n")
    logger.info("run_stats: saved log=%s", path)


def _init_progress(session_id: Optional[str], total: int) -> None:
    if not session_id:
        return
    with PROGRESS_LOCK:
        PROGRESS[session_id] = {"done": 0, "total": total, "status": "running"}


def _update_progress(session_id: Optional[str], done: Optional[int] = None, total: Optional[int] = None,
                     status: Optional[str] = None) -> None:
    if not session_id:
        return
    with PROGRESS_LOCK:
        current = PROGRESS.get(session_id, {"done": 0, "total": 0, "status": "running"})
        if done is not None:
            current["done"] = done
        if total is not None:
            current["total"] = total
        if status is not None:
            current["status"] = status
        PROGRESS[session_id] = current


def _get_progress(session_id: str) -> Dict[str, Any]:
    if not session_id:
        return {"done": 0, "total": 0, "status": "idle"}
    with PROGRESS_LOCK:
        return dict(PROGRESS.get(session_id, {"done": 0, "total": 0, "status": "idle"}))


def _parse_int(value: Any, default: int = 0) -> int:
    if value is None:
        return default
    if isinstance(value, int):
        return value
    text = str(value).strip()
    if not text:
        return default
    try:
        return int(float(text))
    except ValueError:
        return default


def _format_search_result(result: Dict[str, Any]) -> str:
    title = str(result.get("title") or "").strip()
    url = str(result.get("url") or "").strip()
    passages = result.get("passages") or []
    lines: List[str] = []
    if title:
        lines.append(f"Title: {title}")
    if url:
        lines.append(f"URL: {url}")
    if passages:
        lines.append("Snippets:")
        for passage in passages:
            snippet = str(passage).strip()
            if snippet:
                lines.append(f"- {snippet}")
    return "\n".join(lines).strip()


def _safe_filename(value: str, max_len: int = 120) -> str:
    cleaned = []
    for ch in value:
        if ch.isalnum() or ch in ("-", "_"):
            cleaned.append(ch)
        elif ch.isspace() or ch in ("/", "\\", ":", ".", "?", "&", "=", "%"):
            cleaned.append("_")
        else:
            cleaned.append("_")
    safe = "".join(cleaned).strip("_")
    if not safe:
        safe = "log"
    return safe[:max_len]


class State(TypedDict, total=False):
    inputs: List[str]
    tasks: List[str]
    requests: List[str]
    results: List[Dict[str, Any]]
    last_result: int
    pages: List[AnyMessage]
    answers: Dict[str, str]
    messages: List[AnyMessage]
    iterations_left: int
    line: int
    lines: List[Dict[str, Any]]
    input_values: List[str]
    need_more_pages: bool
    done: bool
    session_id: Optional[str]
    total_lines: int
    run_id: Optional[str]


class AnswerItem(BaseModelV1):
    field: str = FieldV1(
        ...,
        alias="Поле",
        description="Название поля из списка tasks, которое нужно заполнить.",
    )
    value: str = FieldV1(
        ...,
        alias="Значение",
        description="Ответ для указанного поля из tasks.",
    )

    class Config:
        allow_population_by_field_name = True


@tool
def return_answers_tool(answers: List[AnswerItem]) -> ToolMessage:
    """Верни ответы списком объектов AnswerItem.

    Формат каждого элемента:
    - Поле: точное название из списка tasks.
    - Значение: найденный ответ (строка).

    Требования:
    - Длина списка = количеству tasks.
    - Для каждого task должен быть один объект.
    - Не выдумывай названия полей и не объединяй несколько задач в одно поле.
    - Если ответа нет, верни пустую строку в Значение.
    """
    payload = [
        {"Поле": item.field, "Значение": item.value}
        for item in answers
        if item.field and item.value
    ]
    return ToolMessage(content=json.dumps(payload, ensure_ascii=False))


@tool
def ask_more_pages() -> ToolMessage:
    """Запросить дополнительные страницы (фрагменты поиска).

    Вызывай, если контекста недостаточно для заполнения всех tasks.
    """
    return ToolMessage(content="MORE_PAGES")


class LLMRequests(BaseModel):
    requests: List[str] = Field(..., min_length=1, max_length=3, description="All necessary browser queries (max 3)")

class ColumnSelection(BaseModel):
    inputs: List[str] = Field(default_factory=list)
    tasks: List[str] = Field(default_factory=list)
    img_task: Optional[str] = None
    img_count: Optional[int] = None


def _filter_columns(values: List[str], columns: List[str], limit: int) -> List[str]:
    seen: set[str] = set()
    filtered: List[str] = []
    for value in values:
        if value in columns and value not in seen:
            filtered.append(value)
            seen.add(value)
        if len(filtered) >= limit:
            break
    return filtered


def _normalize_tasks(values: List[str], limit: int) -> List[str]:
    seen: set[str] = set()
    filtered: List[str] = []
    for value in values:
        name = (value or "").strip()
        if not name or name in seen:
            continue
        filtered.append(name)
        seen.add(name)
        if len(filtered) >= limit:
            break
    return filtered


def select_columns(query: str, columns: List[str], run_id: Optional[str] = None) -> ColumnSelection:
    msgs = [
        SystemMessage(SELECTION_PROMPT),
        HumanMessage(f"Columns: {', '.join(columns)}"),
        HumanMessage(f"User query: {query}"),
    ]
    structured_llm = llm.with_structured_output(ColumnSelection)
    _inc_run_stat(run_id, "llm_count")
    response = structured_llm.invoke(msgs)

    inputs = _filter_columns(response.inputs or [], columns, limit=3)
    tasks = _normalize_tasks(response.tasks or [], limit=3)
    img_task = response.img_task if response.img_task in columns else None
    img_count = response.img_count if isinstance(response.img_count, int) else None
    if img_count is not None and not (1 <= img_count <= 5):
        img_count = None
    if not img_task:
        img_count = None

    selection = ColumnSelection(inputs=inputs, tasks=tasks, img_task=img_task, img_count=img_count)
    logger.info(
        'select_columns: inputs=%s tasks=%s img_task=%s img_count=%s',
        selection.inputs,
        selection.tasks,
        selection.img_task,
        selection.img_count,
    )
    _save_start_agent_log(query, columns, selection)
    return selection


def create_requests_node(state: State) -> State:
    logger.info("create_requests_node: inputs=%s tasks=%s", state.get("inputs", []), state.get("tasks", []))
    inputs = ", ".join(state["inputs"])
    tasks = ", ".join(state["tasks"])
    msgs = [
        SystemMessage(REQUEST_PROMPT),
        HumanMessage(f"Список inputs: {inputs}. Список tasks: {tasks}."),
    ]
    structured_llm = llm.with_structured_output(LLMRequests)
    _inc_run_stat(state.get("run_id"), "llm_count")
    response = structured_llm.invoke(msgs)
    trimmed = (response.requests or [])[:3]
    logger.info("create_requests_node: requests=%s", trimmed)
    return {"requests": trimmed}


def lines_counter_node(state: State) -> State:
    logger.info("lines_counter_node: line=%s total=%s", state.get("line", 0), len(state.get("lines", [])))
    line_idx = int(state.get("line", 0))
    lines = state.get("lines", [])
    session_id = state.get("session_id")
    total_lines = state.get("total_lines") or len(lines)

    if line_idx > 0:
        answers = state.get("answers", {})
        if answers:
            prev_row = lines[line_idx - 1]
            for key, value in answers.items():
                if value:
                    prev_row[key] = value
            logger.info("lines_counter_node: saved answers for line=%s keys=%s", line_idx, list(answers.keys()))
        _update_progress(session_id, done=line_idx, total=total_lines, status="running")

    if line_idx >= len(lines):
        _update_progress(session_id, done=total_lines, total=total_lines, status="done")
        logger.info("lines_counter_node: done")
        return {"done": True}

    current = lines[line_idx]
    input_values = [str(current.get(col, "")).strip() for col in state["inputs"]]
    row_context = " | ".join(f"{col}: {current.get(col, '')}" for col in state["inputs"])
    task_context = ", ".join(state["tasks"])

    messages: List[AnyMessage] = [
        SystemMessage(AGENT_PROMPT),
        SystemMessage("????? ?????? ????? return_answers_tool ??????? ???????? AnswerItem. ?????? ?????? = ?????????? tasks. ? ???? ????? ?????? ??? ?? tasks, ? ???????? ? ????????? ????? (??????)."),
        HumanMessage(f"Текущая строка: {row_context}. Нужно заполнить: {task_context}."),    ]

    return {
        "line": line_idx + 1,
        "input_values": input_values,
        "results": [],
        "pages": [],
        "last_result": 0,
        "answers": {},
        "messages": messages,
        "iterations_left": MAX_ITERATIONS,
        "need_more_pages": False,
        "done": False,
    }


def get_urls_node(state: State) -> State:
    logger.info("get_urls_node: requests=%s", state.get("requests", []))
    requests_templates = state.get("requests", [])
    input_values = state.get("input_values", [])
    results: List[Dict[str, Any]] = []
    lr_value: Optional[int] = None
    if XML_SEARCH_LR:
        try:
            lr_value = int(XML_SEARCH_LR)
        except ValueError:
            logger.warning("xml_search: invalid XML_SEARCH_LR=%r", XML_SEARCH_LR)

    for template in requests_templates:
        request = template
        for idx, value in enumerate(input_values, start=1):
            request = request.replace(f"*input_{idx}*", value)
        _inc_run_stat(state.get("run_id"), "search_count")
        try:
            query_results = xml_search(
                request,
                docs=XML_SEARCH_DOCS,
                maxpassages=XML_SEARCH_MAXPASSAGES,
                lr=lr_value,
                groupby=XML_SEARCH_GROUPBY,
            )
        except Exception as exc:
            logger.warning("xml_search failed query=%r error=%s", request, exc)
            continue
        logger.info("xml_search: results=%d query=%r", len(query_results), request)
        _save_xml_search_log(request, query_results)
        results.extend(query_results)

    logger.info("get_urls_node: results=%s", len(results))
    return {"results": results}


def parse_pages_node(state: State) -> State:
    logger.info(
        "parse_pages_node: last_result=%s total_results=%s",
        state.get("last_result", 0),
        len(state.get("results", [])),
    )
    last_result = int(state.get("last_result", 0))
    results = state.get("results", [])
    pages = list(state.get("pages", []))
    messages = list(state.get("messages", []))

    if last_result >= len(results):
        logger.info("parse_pages_node: no more results")
        return {"need_more_pages": False}

    end = min(last_result + PAGES_PER_TIME, len(results))
    for idx in range(last_result, end):
        result = results[idx]
        content = _format_search_result(result)
        if not content:
            continue
        pages.append(SystemMessage(content=content))
        messages.append(SystemMessage(content=content))

    logger.info("parse_pages_node: parsed=%s new_last_result=%s", end - last_result, end)
    return {
        "pages": pages,
        "messages": messages,
        "last_result": end,
        "need_more_pages": False,
    }


def _extract_tool_calls(message: AIMessage) -> List[Dict[str, Any]]:
    tool_calls = getattr(message, "tool_calls", None)
    if tool_calls:
        return list(tool_calls)
    extra = getattr(message, "additional_kwargs", {}) or {}
    if isinstance(extra, dict):
        calls = extra.get("tool_calls")
        if calls:
            return list(calls)
    return []


def _normalize_tool_call(call: Dict[str, Any]) -> Dict[str, Any]:
    if "function" in call and isinstance(call["function"], dict):
        func = call["function"]
        name = func.get("name")
        args = func.get("arguments")
        if isinstance(args, str):
            try:
                args = json.loads(args)
            except ValueError:
                args = {}
        return {
            "id": call.get("id"),
            "name": name,
            "args": args if isinstance(args, dict) else {},
        }
    return {
        "id": call.get("id") or call.get("tool_call_id"),
        "name": call.get("name") or call.get("tool"),
        "args": call.get("args") if isinstance(call.get("args"), dict) else {},
    }


def agent_node(state: State) -> State:
    logger.info("agent_node: iterations_left=%s", state.get("iterations_left", 0))
    llm_with_tools = llm.bind_tools([return_answers_tool, ask_more_pages])
    msgs: List[AnyMessage] = list(state.get("messages", []))

    if int(state.get("iterations_left", 0)) <= 1:
        msgs.append(SystemMessage(content=LAST_PROMPT))

    _inc_run_stat(state.get("run_id"), "llm_count")
    response = None
    for attempt in range(1, LLM_RETRIES + 1):
        try:
            response = llm_with_tools.invoke(msgs)
            break
        except Exception as exc:
            logger.warning(
                "agent_node: llm invoke failed attempt=%s/%s error=%s",
                attempt,
                LLM_RETRIES,
                exc,
            )
            if attempt < LLM_RETRIES:
                time.sleep(LLM_RETRY_WAIT)
    if response is None:
        logger.error("agent_node: llm invoke failed after retries; skipping line")
        return {
            "messages": msgs,
            "iterations_left": 0,
            "need_more_pages": False,
        }
    msgs.append(response)

    next_state: State = {
        "messages": msgs,
        "iterations_left": int(state.get("iterations_left", 0)) - 1,
        "need_more_pages": False,
    }

    tool_calls = _extract_tool_calls(response)
    if not tool_calls:
        logger.info("agent_node: no tool calls")
        next_state["need_more_pages"] = False
        next_state["iterations_left"] = 0
        return next_state

    for call in tool_calls:
        normalized = _normalize_tool_call(call)
        name = normalized.get("name")
        args = normalized.get("args") or {}
        call_id = normalized.get("id") or "manual"
        if name == "return_answers_tool":
            raw_answers: Any = args.get("answers") if isinstance(args, dict) else None
            items: List[Dict[str, Any]] = []
            if isinstance(raw_answers, list):
                items = [item for item in raw_answers if isinstance(item, dict)]
            elif isinstance(raw_answers, dict):
                items = [raw_answers]
            elif isinstance(args, list):
                items = [item for item in args if isinstance(item, dict)]
            elif isinstance(args, dict):
                items = [args]

            answers_map: Dict[str, str] = {}
            for item in items:
                field = item.get("Поле") or item.get("field")
                value = item.get("Значение") or item.get("value")
                if field and value:
                    answers_map[str(field).strip()] = str(value).strip()

            cleaned: Dict[str, str] = {}
            for key in state.get("tasks", []):
                value = answers_map.get(key)
                if value:
                    cleaned[key] = value

            if cleaned:
                next_state["answers"] = cleaned
            else:
                next_state["iterations_left"] = 0
            next_state["messages"].append(
                ToolMessage(content=json.dumps(cleaned, ensure_ascii=False), tool_call_id=call_id)
            )
            logger.info("agent_node: return_answers_tool keys=%s", list(cleaned.keys()))
            continue
        if name == "ask_more_pages":
            next_state["need_more_pages"] = True
            next_state["messages"].append(ToolMessage(content="MORE_PAGES", tool_call_id=call_id))
            logger.info("agent_node: ask_more_pages")
            continue
        next_state["messages"].append(ToolMessage(content="", tool_call_id=call_id))
        logger.warning("agent_node: unknown tool call name=%s", name)

    return next_state


def _router(state: State) -> Literal["parse_pages", "line_counter", "agent", "end"]:
    if state.get("done"):
        logger.info("_router: done")
        return "end"
    if state.get("answers") or int(state.get("iterations_left", 0)) <= 0:
        logger.info("_router: line_counter")
        return "line_counter"
    if state.get("need_more_pages"):
        logger.info("_router: parse_pages")
        return "parse_pages"
    logger.info("_router: agent")
    return "agent"


def _line_router(state: State) -> Literal["get_urls", "end"]:
    if state.get("done"):
        logger.info("_line_router: done")
        return "end"
    logger.info("_line_router: get_urls")
    return "get_urls"


sub_workflow = StateGraph(State)
sub_workflow.add_node("create_requests", create_requests_node)
sub_workflow.add_node("line_counter", lines_counter_node)
sub_workflow.add_node("get_urls", get_urls_node)
sub_workflow.add_node("parse_pages", parse_pages_node)
sub_workflow.add_node("agent", agent_node)

sub_workflow.set_entry_point("create_requests")
sub_workflow.add_edge("create_requests", "line_counter")
sub_workflow.add_conditional_edges("line_counter", _line_router, {"get_urls": "get_urls", "end": END})
sub_workflow.add_edge("get_urls", "parse_pages")
sub_workflow.add_edge("parse_pages", "agent")
sub_workflow.add_conditional_edges(
    "agent",
    _router,
    {"parse_pages": "parse_pages", "line_counter": "line_counter", "agent": "agent", "end": END},
)

sub_graph = sub_workflow.compile()


def _parse_list(value: str) -> List[str]:
    try:
        data = json.loads(value)
    except json.JSONDecodeError:
        data = None
    if isinstance(data, list):
        return [str(item).strip() for item in data if str(item).strip()]
    return [item.strip() for item in value.split(",") if item.strip()]


def _is_row_empty(row: Dict[str, Any]) -> bool:
    for value in row.values():
        if value is None:
            continue
        if isinstance(value, str):
            if value.strip() == "":
                continue
            return False
        return False
    return True


def _ensure_image_columns(columns: List[str]) -> None:
    if "image_1" in columns:
        return
    for idx in range(1, 6):
        col = f"image_{idx}"
        if col not in columns:
            columns.append(col)


def _load_rows_from_csv(data: bytes) -> List[Dict[str, Any]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [row for row in reader if not _is_row_empty(row)]


def _load_rows_from_xlsx(data: bytes) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=400, detail="openpyxl не установлен для чтения xlsx") from exc

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    results: List[Dict[str, Any]] = []
    for row in rows[1:]:
        row_dict: Dict[str, Any] = {}
        for idx, header in enumerate(headers):
            row_dict[header] = row[idx] if idx < len(row) else None
        if not _is_row_empty(row_dict):
            results.append(row_dict)
    return results


def _write_rows_to_csv(rows: List[Dict[str, Any]], fieldnames: List[str]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def _write_rows_to_xlsx(rows: List[Dict[str, Any]], fieldnames: List[str]) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise HTTPException(status_code=400, detail="openpyxl не установлен для записи xlsx") from exc

    wb = Workbook()
    ws = wb.active
    ws.append(fieldnames)
    for row in rows:
        ws.append([row.get(name) for name in fieldnames])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def process_rows(
    lines: List[Dict[str, Any]],
    inputs: List[str],
    tasks: List[str],
    session_id: Optional[str] = None,
    total_lines: Optional[int] = None,
    run_id: Optional[str] = None,
) -> List[Dict[str, Any]]:
    initial_state: State = {
        "inputs": inputs,
        "tasks": tasks,
        "lines": lines,
        "line": 0,
        "iterations_left": MAX_ITERATIONS,
        "messages": [],
        "session_id": session_id,
        "total_lines": total_lines or len(lines),
        "run_id": run_id,
    }
    final_state = sub_graph.invoke(initial_state, config={"recursion_limit": RECURSION_LIMIT})
    return final_state.get("lines", lines)


async def process_rows_async(
    lines: List[Dict[str, Any]],
    inputs: List[str],
    tasks: List[str],
    session_id: Optional[str] = None,
    total_lines: Optional[int] = None,
    run_id: Optional[str] = None,
    image_search_column: Optional[str] = None,
    image_count_value: Optional[int] = None,
    columns: Optional[List[str]] = None,
) -> List[Dict[str, Any]]:
    semaphore = asyncio.Semaphore(MAX_CONCURRENCY)
    results: List[Optional[Dict[str, Any]]] = [None] * len(lines)
    progress_lock = asyncio.Lock()
    completed = 0
    total = total_lines or len(lines)

    async def handle_row(idx: int, row: Dict[str, Any]) -> None:
        nonlocal completed
        async with semaphore:
            row_copy = dict(row)
            processed = await asyncio.to_thread(
                process_rows, [row_copy], inputs, tasks, None, 1, run_id
            )
            processed_row = processed[0] if processed else row_copy
            if image_search_column and image_count_value and columns is not None:
                await asyncio.to_thread(
                    process_images,
                    [processed_row],
                    columns,
                    image_search_column,
                    "",
                    image_count_value,
                    run_id,
                )
            results[idx] = processed_row
        if session_id:
            async with progress_lock:
                completed += 1
                _update_progress(session_id, done=completed, total=total, status="running")

    await asyncio.gather(*(handle_row(idx, row) for idx, row in enumerate(lines)))
    return [results[idx] if results[idx] is not None else lines[idx] for idx in range(len(lines))]


async def process_images_async(
    rows: List[Dict[str, Any]],
    columns: List[str],
    image_search_column: str,
    image_count_value: Optional[int],
    session_id: Optional[str] = None,
    run_id: Optional[str] = None,
) -> List[Dict[str, Any]]:
    semaphore = asyncio.Semaphore(MAX_CONCURRENCY)
    progress_lock = asyncio.Lock()
    completed = 0
    total = len(rows)

    async def handle_row(idx: int, row: Dict[str, Any]) -> None:
        nonlocal completed
        async with semaphore:
            await asyncio.to_thread(
                process_images,
                [row],
                columns,
                image_search_column,
                "",
                image_count_value,
                run_id,
            )
        if session_id:
            async with progress_lock:
                completed += 1
                _update_progress(session_id, done=completed, total=total, status="running")

    await asyncio.gather(*(handle_row(idx, row) for idx, row in enumerate(rows)))
    return rows


def process_rows_with_images_sync(
    lines: List[Dict[str, Any]],
    inputs: List[str],
    tasks: List[str],
    session_id: Optional[str],
    total_lines: Optional[int],
    run_id: Optional[str],
    image_search_column: Optional[str],
    image_count_value: Optional[int],
    columns: List[str],
) -> List[Dict[str, Any]]:
    results: List[Dict[str, Any]] = []
    completed = 0
    total = total_lines or len(lines)
    for row in lines:
        row_copy = dict(row)
        processed = process_rows([row_copy], inputs, tasks, None, 1, run_id)
        processed_row = processed[0] if processed else row_copy
        if image_search_column and image_count_value:
            process_images(
                [processed_row],
                columns,
                image_search_column,
                "",
                image_count_value,
                run_id,
            )
        results.append(processed_row)
        if session_id:
            completed += 1
            _update_progress(session_id, done=completed, total=total, status="running")
    return results


def process_images_sync(
    rows: List[Dict[str, Any]],
    columns: List[str],
    image_search_column: str,
    image_count_value: Optional[int],
    session_id: Optional[str],
    run_id: Optional[str],
) -> List[Dict[str, Any]]:
    completed = 0
    total = len(rows)
    for row in rows:
        process_images(
            [row],
            columns,
            image_search_column,
            "",
            image_count_value,
            run_id,
        )
        if session_id:
            completed += 1
            _update_progress(session_id, done=completed, total=total, status="running")
    return rows


def process_payload(
    data: bytes, filename: str, query: str, session_id: Optional[str]
) -> tuple[bytes, str, str]:
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    _init_run_stats(run_id)
    no_limit = False
    stripped_query = query.lstrip()
    if stripped_query.startswith("#"):
        no_limit = True
        query = stripped_query[1:].lstrip()

    ext = os.path.splitext(filename)[1].lower()
    if ext in (".csv", ""):
        rows = _load_rows_from_csv(data)
        fmt = "csv"
    elif ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        rows = _load_rows_from_xlsx(data)
        fmt = "xlsx"
    else:
        raise HTTPException(status_code=400, detail=f"unsupported file type: {ext}")

    if not rows:
        raise HTTPException(status_code=400, detail="file has no data")

    if not no_limit and len(rows) > MAX_ROWS:
        rows = rows[:MAX_ROWS]

    _init_progress(session_id, total=len(rows))

    try:
        columns = list(rows[0].keys())

        selection = select_columns(query, columns, run_id)
        inputs_list = selection.inputs
        tasks_list = selection.tasks
        image_search_column = selection.img_task
        image_count_value = selection.img_count

        if tasks_list and not inputs_list:
            raise HTTPException(status_code=400, detail="LLM did not select inputs for tasks")

        for name in tasks_list:
            if name not in columns:
                columns.append(name)

        if image_search_column and image_count_value:
            _ensure_image_columns(columns)

        processed_rows = rows
        if tasks_list:
            try:
                processed_rows = asyncio.run(
                    process_rows_async(
                        rows,
                        inputs_list,
                        tasks_list,
                        session_id,
                        len(rows),
                        run_id,
                        image_search_column,
                        image_count_value,
                        columns,
                    )
                )
            except RuntimeError:
                processed_rows = process_rows_with_images_sync(
                    rows,
                    inputs_list,
                    tasks_list,
                    session_id,
                    len(rows),
                    run_id,
                    image_search_column,
                    image_count_value,
                    columns,
                )
        elif image_search_column and image_count_value:
            try:
                processed_rows = asyncio.run(
                    process_images_async(
                        processed_rows,
                        columns,
                        image_search_column,
                        image_count_value,
                        session_id,
                        run_id,
                    )
                )
            except RuntimeError:
                processed_rows = process_images_sync(
                    processed_rows,
                    columns,
                    image_search_column,
                    image_count_value,
                    session_id,
                    run_id,
                )
        if not tasks_list and not image_search_column:
            _update_progress(session_id, done=len(rows), total=len(rows), status="done")

        if fmt == "csv":
            payload = _write_rows_to_csv(processed_rows, columns)
            media = "text/csv"
            out_name = os.path.splitext(filename)[0] + "_filled.csv"
        else:
            payload = _write_rows_to_xlsx(processed_rows, columns)
            media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            out_name = os.path.splitext(filename)[0] + "_filled.xlsx"

        _update_progress(session_id, done=len(rows), total=len(rows), status="done")

        return payload, media, out_name
    finally:
        _write_run_stats(run_id)




def process_images(
    rows: List[Dict[str, Any]],
    columns: List[str],
    search_column: str,
    count_column: str,
    count_value: Optional[int],
    run_id: Optional[str] = None,
) -> List[Dict[str, Any]]:
    if not search_column or (not count_column and count_value is None):
        return rows

    if "image_1" not in columns:
        for idx in range(1, 6):
            col = f"image_{idx}"
            if col not in columns:
                columns.append(col)

    lr_value: Optional[int] = None
    if XML_IMAGES_LR:
        try:
            lr_value = int(XML_IMAGES_LR)
        except ValueError:
            logger.warning("xml_images: invalid XML_IMAGES_LR=%r", XML_IMAGES_LR)

    for row in rows:
        query = str(row.get(search_column, "")).strip()
        if not query:
            continue
        if count_value is not None:
            target = count_value
        else:
            target = _parse_int(row.get(count_column), default=0)
        if target <= 0:
            continue
        if target > 5:
            target = 5

        _inc_run_stat(run_id, "image_search_count")
        try:
            urls = xml_images(
                query,
                limit=XML_IMAGES_LIMIT,
                lr=lr_value,
                page=XML_IMAGES_PAGE,
                device=XML_IMAGES_DEVICE,
                domain=XML_IMAGES_DOMAIN,
            )
        except Exception as exc:
            logger.warning("xml_images failed query=%r error=%s", query, exc)
            continue

        _save_image_search_log(query, urls)
        found: List[str] = []
        for url in urls:
            try:
                if has_watermark_url(url):
                    logger.info("xml_images: watermark found url=%s", url)
                    continue
            except Exception as exc:
                logger.warning("xml_images: watermark check failed url=%s error=%s", url, exc)
                continue
            found.append(url)
            if len(found) >= target:
                break

        for idx in range(1, target + 1):
            row[f"image_{idx}"] = found[idx - 1] if idx - 1 < len(found) else ""

    return rows


app = FastAPI()


class SessionStartResponse(BaseModel):
    session_id: str


class SessionEndRequest(BaseModel):
    session_id: str


class ProgressResponse(BaseModel):
    done: int
    total: int
    status: str


class SurveyCreate(BaseModel):
    session_id: str
    purpose: str
    additions: str
    contact: str


def _store_survey(payload: SurveyCreate) -> None:
    with _get_db_conn() as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO surveys (session_id, purpose, additions, contact, created_at)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (payload.session_id, payload.purpose, payload.additions, payload.contact, _utc_now()),
        )


@app.post("/api/session/start", response_model=SessionStartResponse)
async def session_start(request: Request) -> SessionStartResponse:
    ip_hash = _hash_ip(_get_client_ip(request))
    user_agent = request.headers.get("user-agent")
    referrer = request.headers.get("referer")
    visitor_id = await asyncio.to_thread(_get_or_create_visitor, ip_hash)
    session_id = await asyncio.to_thread(_create_session, visitor_id, user_agent, referrer)
    return SessionStartResponse(session_id=session_id)


@app.post("/api/session/end")
async def session_end(payload: SessionEndRequest) -> Dict[str, str]:
    await asyncio.to_thread(_end_session, payload.session_id)
    return {"status": "ok"}


@app.get("/api/progress", response_model=ProgressResponse)
async def get_progress(session_id: str = Query(...)) -> ProgressResponse:
    data = _get_progress(session_id)
    return ProgressResponse(
        done=int(data.get("done", 0)),
        total=int(data.get("total", 0)),
        status=str(data.get("status", "idle")),
    )


@app.post("/api/survey")
async def save_survey(payload: SurveyCreate) -> Dict[str, str]:
    await asyncio.to_thread(_store_survey, payload)
    return {"status": "ok"}


@app.post("/")
async def process_file(
    request: Request,
    file: UploadFile = File(...),
    query: str = Form(...),
    session_id: Optional[str] = Form(None),
) -> StreamingResponse:
    query = query.strip()
    if not query:
        raise HTTPException(status_code=400, detail="query is required")

    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="file is empty")

    filename = file.filename or "input.csv"

    ip_hash = _hash_ip(_get_client_ip(request))
    user_agent = request.headers.get("user-agent")
    referrer = request.headers.get("referer")
    visitor_id = await asyncio.to_thread(_get_or_create_visitor, ip_hash)
    if not session_id:
        session_id = await asyncio.to_thread(_create_session, visitor_id, user_agent, referrer)

    await asyncio.to_thread(
        _store_upload,
        session_id,
        filename,
        file.content_type or "application/octet-stream",
        data,
    )

    payload, media, out_name = await asyncio.to_thread(
        process_payload, data, filename, query, session_id
    )

    # Estimate row count from progress state (set by process_payload)
    progress_data = _get_progress(session_id)
    rows_processed = int(progress_data.get("total", 0))

    # Save result to disk and record in DB (best-effort, never blocks response)
    try:
        await asyncio.to_thread(
            _store_result,
            session_id,
            filename,
            out_name,
            payload,
            query,
            rows_processed,
        )
    except Exception as exc:
        logger.warning("Failed to store result: %s", exc)

    safe_name = urllib.parse.quote(out_name)
    return StreamingResponse(
        io.BytesIO(payload),
        media_type=media,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}"},
    )


@app.post("/process")
async def process_file_alias(
    request: Request,
    file: UploadFile = File(...),
    query: str = Form(...),
    session_id: Optional[str] = Form(None),
) -> Dict[str, str]:
    return await process_file(
        request=request,
        file=file,
        query=query,
        session_id=session_id,
    )


@app.post("/api/process")
async def process_file_api_alias(
    request: Request,
    file: UploadFile = File(...),
    query: str = Form(...),
    session_id: Optional[str] = Form(None),
) -> Dict[str, str]:
    return await process_file(
        request=request,
        file=file,
        query=query,
        session_id=session_id,
    )


# ─────────────────────────────────────────────────────────────────────────────
# ADMIN API
# ─────────────────────────────────────────────────────────────────────────────

ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "changeme")
ADMIN_SECRET = os.getenv("ADMIN_SECRET", "stroki_admin_secret_key_change_me")
_ADMIN_TOKEN: Optional[str] = None
_ADMIN_TOKEN_LOCK = Lock()


def _make_admin_token() -> str:
    payload = f"{ADMIN_USERNAME}:{time.time()}"
    sig = hmac.new(ADMIN_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return f"{payload}:{sig}"


def _verify_admin_token(token: str) -> bool:
    import hmac
    global _ADMIN_TOKEN
    with _ADMIN_TOKEN_LOCK:
        return token == _ADMIN_TOKEN


def _require_admin(request: Request) -> None:
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Unauthorized")
    token = auth[len("Bearer "):]
    if not _verify_admin_token(token):
        raise HTTPException(status_code=401, detail="Invalid or expired token")


class AdminLoginRequest(BaseModel):
    username: str
    password: str


class AdminLoginResponse(BaseModel):
    token: str


@app.post("/api/admin/login", response_model=AdminLoginResponse)
async def admin_login(payload: AdminLoginRequest) -> AdminLoginResponse:
    global _ADMIN_TOKEN
    if payload.username != ADMIN_USERNAME or payload.password != ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Неверный логин или пароль.")
    token = _make_admin_token()
    with _ADMIN_TOKEN_LOCK:
        _ADMIN_TOKEN = token
    return AdminLoginResponse(token=token)


@app.get("/api/admin/me")
async def admin_me(request: Request) -> Dict[str, str]:
    _require_admin(request)
    return {"username": ADMIN_USERNAME}


@app.get("/api/admin/stats")
async def admin_stats(request: Request) -> Dict[str, Any]:
    _require_admin(request)

    def _fetch() -> Dict[str, Any]:
        with _get_db_conn() as conn, conn.cursor() as cur:
            today = datetime.now(timezone.utc).date().isoformat()

            cur.execute("SELECT COUNT(*) FROM visitors")
            total_visitors = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM sessions")
            total_sessions = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM sessions WHERE started_at::date = %s", (today,))
            sessions_today = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM uploads")
            total_uploads = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM uploads WHERE created_at::date = %s", (today,))
            uploads_today = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM results")
            total_results = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM results WHERE created_at::date = %s", (today,))
            results_today = cur.fetchone()[0]

            cur.execute("SELECT COUNT(*) FROM surveys")
            total_surveys = cur.fetchone()[0]

            return {
                "total_visitors": total_visitors,
                "total_sessions": total_sessions,
                "sessions_today": sessions_today,
                "total_uploads": total_uploads,
                "uploads_today": uploads_today,
                "total_results": total_results,
                "results_today": results_today,
                "total_surveys": total_surveys,
            }

    return await asyncio.to_thread(_fetch)


@app.get("/api/admin/sessions")
async def admin_sessions(request: Request, limit: int = 200) -> List[Dict[str, Any]]:
    _require_admin(request)

    def _fetch() -> List[Dict[str, Any]]:
        with _get_db_conn() as conn, conn.cursor() as cur:
            cur.execute(
                """
                SELECT s.id, s.started_at, s.ended_at, s.user_agent, s.referrer,
                       COUNT(DISTINCT u.id) AS uploads_count,
                       COUNT(DISTINCT r.id) AS results_count
                FROM sessions s
                LEFT JOIN uploads u ON u.session_id = s.id
                LEFT JOIN results r ON r.session_id = s.id
                GROUP BY s.id
                ORDER BY s.started_at DESC
                LIMIT %s
                """,
                (limit,),
            )
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            result = []
            for row in rows:
                d = dict(zip(cols, row))
                for k in ("started_at", "ended_at"):
                    if d[k] is not None:
                        d[k] = d[k].isoformat()
                result.append(d)
            return result

    return await asyncio.to_thread(_fetch)


@app.get("/api/admin/uploads")
async def admin_uploads(request: Request, limit: int = 300) -> List[Dict[str, Any]]:
    _require_admin(request)

    def _fetch() -> List[Dict[str, Any]]:
        with _get_db_conn() as conn, conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, session_id, original_name, stored_name, mime, size_bytes, sha256, created_at
                FROM uploads
                ORDER BY created_at DESC
                LIMIT %s
                """,
                (limit,),
            )
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            result = []
            for row in rows:
                d = dict(zip(cols, row))
                if d["created_at"] is not None:
                    d["created_at"] = d["created_at"].isoformat()
                d["id"] = str(d["id"])
                d["session_id"] = str(d["session_id"])
                result.append(d)
            return result

    return await asyncio.to_thread(_fetch)


@app.get("/api/admin/results")
async def admin_results(request: Request, limit: int = 300) -> List[Dict[str, Any]]:
    _require_admin(request)

    def _fetch() -> List[Dict[str, Any]]:
        with _get_db_conn() as conn, conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, session_id, upload_original_name, result_name, stored_name,
                       size_bytes, query, rows_processed, created_at
                FROM results
                ORDER BY created_at DESC
                LIMIT %s
                """,
                (limit,),
            )
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            result = []
            for row in rows:
                d = dict(zip(cols, row))
                if d["created_at"] is not None:
                    d["created_at"] = d["created_at"].isoformat()
                d["id"] = str(d["id"])
                d["session_id"] = str(d["session_id"])
                result.append(d)
            return result

    return await asyncio.to_thread(_fetch)


@app.get("/api/admin/results/{result_id}/download")
async def admin_result_download(result_id: str, request: Request) -> StreamingResponse:
    _require_admin(request)

    def _fetch():
        with _get_db_conn() as conn, conn.cursor() as cur:
            cur.execute(
                "SELECT stored_name, result_name FROM results WHERE id = %s",
                (result_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="Result not found")
            return row

    stored_name, result_name = await asyncio.to_thread(_fetch)
    path = os.path.join(UPLOAD_DIR, stored_name)
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="File not found on disk")

    with open(path, "rb") as f:
        data = f.read()

    ext = os.path.splitext(result_name)[1].lower()
    if ext == ".csv":
        media = "text/csv"
    else:
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    safe_name = urllib.parse.quote(result_name)
    return StreamingResponse(
        io.BytesIO(data),
        media_type=media,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}"},
    )


@app.get("/api/admin/surveys")
async def admin_surveys(request: Request, limit: int = 300) -> List[Dict[str, Any]]:
    _require_admin(request)

    def _fetch() -> List[Dict[str, Any]]:
        with _get_db_conn() as conn, conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, session_id, purpose, additions, contact, created_at
                FROM surveys
                ORDER BY created_at DESC
                LIMIT %s
                """,
                (limit,),
            )
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            result = []
            for row in rows:
                d = dict(zip(cols, row))
                if d["created_at"] is not None:
                    d["created_at"] = d["created_at"].isoformat()
                d["id"] = str(d["id"])
                d["session_id"] = str(d["session_id"])
                result.append(d)
            return result

    return await asyncio.to_thread(_fetch)


# ─── Helper: save result file to disk and record in DB ───────────────────────

def _store_result(
    session_id: str,
    upload_original_name: str,
    result_name: str,
    data: bytes,
    query: str,
    rows_processed: int,
    upload_id: Optional[str] = None,
) -> None:
    """Save processed result to disk and insert a record into the results table."""
    ext = os.path.splitext(result_name)[1].lower() or ".csv"
    date_dir = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    stored_dir = os.path.join(UPLOAD_DIR, date_dir)
    os.makedirs(stored_dir, exist_ok=True)
    stored_name = f"result_{uuid.uuid4().hex}{ext}"
    stored_path = os.path.join(stored_dir, stored_name)

    with open(stored_path, "wb") as f:
        f.write(data)

    relative_name = os.path.join(date_dir, stored_name).replace("\\", "/")
    with _get_db_conn() as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO results
              (session_id, upload_id, upload_original_name, result_name, stored_name,
               size_bytes, query, rows_processed, created_at)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                session_id,
                upload_id,
                upload_original_name,
                result_name,
                relative_name,
                len(data),
                query,
                rows_processed,
                _utc_now(),
            ),
        )


